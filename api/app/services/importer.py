"""Bulk inventory import — Python port of web/lib/importer.ts (Method A & B)."""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import datetime, timezone
from types import SimpleNamespace

from ..catalog import makes, part_types
from .query_parser import duplicate_key, normalize, normalize_oem, token_matches

TEMPLATE_COLUMNS = [
    "sku", "part_name", "part_type", "oem_number", "make", "model", "generation", "year_from", "year_to",
    "engine_code", "transmission", "market_spec", "condition", "price_aed", "quantity", "warranty_days",
    "description", "donor_vehicle", "status",
]

CONDITIONS = {"used": "used", "مستعمل": "used", "tested": "tested", "مفحوص": "tested", "refurbished": "refurbished", "rebuilt": "refurbished", "مجدد": "refurbished"}


def resolve_make(v: str) -> str | None:
    n = normalize(v or "")
    if not n:
        return None
    for m in makes():
        names = [normalize(x) for x in [m["id"], m["name"], m["nameAr"], *m["aliases"]]]
        if n in names or any(token_matches(n, x) for x in names):
            return m["id"]
    return None


def resolve_model(make_id: str, v: str) -> str | None:
    n = normalize(v or "")
    mk = next((m for m in makes() if m["id"] == make_id), None)
    if not mk or not n:
        return None
    for md in mk["models"]:
        names = [normalize(x) for x in [md["id"], md["name"], md["nameAr"], md["id"].replace("-", " "), *md.get("aliases", [])]]
        if n in names or any(token_matches(n, x) for x in names):
            return md["id"]
        if n in [normalize(v["name"]) for g in md["generations"] for v in g["variants"]]:
            return md["id"]
    return None


def resolve_part_type(v: str) -> dict | None:
    n = normalize(v or "")
    if not n:
        return None
    types = list(part_types().values())
    return (
        next((t for t in types if t["id"] == n or normalize(t["en"]) == n or normalize(t["ar"]) == n), None)
        or next((t for t in types if any(normalize(s) == n for s in t["synonyms"])), None)
        or next((t for t in types if any(token_matches(n, normalize(s)) for s in t["synonyms"])), None)
    )


def _int(v, default=None):
    try:
        return int(float(str(v).replace(",", ""))) if str(v).strip() != "" else default
    except ValueError:
        return None


@dataclass
class ImportPlan:
    create: list[dict] = field(default_factory=list)
    update: list[tuple[str, dict]] = field(default_factory=list)
    duplicates: int = 0
    errors: list[dict] = field(default_factory=list)
    total: int = 0


def plan_import(rows: list[dict], *, supplier_id: str, emirate: str, area: str, existing: list, source: str = "csv") -> ImportPlan:
    now = datetime.now(timezone.utc)
    plan = ImportPlan(total=len(rows))
    by_sku = {p.sku: p for p in existing if p.sku}
    seen_keys = {duplicate_key(p) for p in existing if p.status != "sold"}
    seen_skus: set[str] = set()

    for idx, raw in enumerate(rows):
        row_no = idx + 2
        row = {str(k).strip().lower().replace(" ", "_"): ("" if v is None else str(v).strip()) for k, v in raw.items() if k is not None}
        if not any(row.values()):
            plan.total -= 1
            continue
        errs: list[str] = []
        make = resolve_make(row.get("make", ""))
        if not make:
            errs.append(f'Unknown make "{row.get("make", "")}"')
        model = resolve_model(make, row.get("model", "")) if make else None
        if make and not model:
            errs.append(f'Unknown model "{row.get("model", "")}" for {make}')
        pt = resolve_part_type(row.get("part_type") or row.get("part_name", ""))
        if not pt:
            errs.append(f'Could not map part type "{row.get("part_type") or row.get("part_name", "")}" to the taxonomy')
        year_from = _int(row.get("year_from"))
        year_to = _int(row.get("year_to") or row.get("year_from"))
        if not (year_from and 1980 <= year_from <= 2030):
            errs.append("year_from must be between 1980 and 2030")
        if not (year_to and year_from and year_from <= year_to <= 2030):
            errs.append("year_to must be ≥ year_from")
        try:
            price = float(row.get("price_aed", "").replace(",", ""))
        except ValueError:
            price = -1
        if not (0 < price < 1_000_000):
            errs.append("price_aed must be a positive number")
        quantity = _int(row.get("quantity"), 1)
        if quantity is None or not (0 <= quantity < 10000):
            errs.append("quantity must be a whole number")
        condition = CONDITIONS.get(normalize(row.get("condition") or "used"))
        if not condition:
            errs.append("condition must be used, tested or refurbished")
        warranty = _int(row.get("warranty_days"), 0)
        if warranty is None or not (0 <= warranty <= 730):
            errs.append("warranty_days must be 0–730")
        status = row.get("status") if row.get("status") in ("available", "reserved", "sold") else "available"
        if pt and pt.get("safetyCritical") and not row.get("oem_number"):
            errs.append(f'{pt["en"]} is safety-critical: oem_number is required')
        sku = row.get("sku") or None
        if sku and sku in seen_skus:
            errs.append(f'Duplicate sku "{sku}" inside the file')
        if errs:
            plan.errors.append({"row": row_no, "message": "; ".join(errs)})
            continue
        if sku:
            seen_skus.add(sku)
        spec = (row.get("market_spec") or "").upper()
        record = dict(
            supplier_id=supplier_id, sku=sku, name=row.get("part_name") or pt["en"], category_id=pt["categoryId"], part_type_id=pt["id"],
            oem=normalize_oem(row["oem_number"]) if row.get("oem_number") else None, make=make, model=model,
            generation=row.get("generation") or None, year_from=year_from, year_to=year_to, engine=row.get("engine_code") or None,
            transmission=row.get("transmission") or None, market_spec="GCC" if spec == "GCC" else ("Imported" if spec else "Unknown"),
            condition=condition, price=round(price), quantity=quantity, warranty_days=warranty, emirate=emirate, area=area,
            description=row.get("description", ""), donor_vehicle=row.get("donor_vehicle") or None,
            status="sold" if quantity == 0 else status, last_verified=now,
            fitment="verified" if row.get("oem_number") else "likely",
            moderation="pending" if pt.get("safetyCritical") else "approved", source=source, is_sample=False,
        )
        if sku and sku in by_sku:
            plan.update.append((by_sku[sku].id, record))
            continue
        key = duplicate_key(SimpleNamespace(**record))
        if key in seen_keys:
            plan.duplicates += 1
            plan.errors.append({"row": row_no, "message": "Duplicate of an existing listing (same part type and OEM / fitment) — skipped"})
            continue
        seen_keys.add(key)
        plan.create.append(record)
    return plan


def read_upload(filename: str, data: bytes) -> list[dict]:
    """Parse CSV (UTF-8, with or without BOM) or XLSX into row dicts."""
    name = filename.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h or "").strip() for h in rows[0]]
        return [{header[i]: ("" if v is None else v) for i, v in enumerate(r) if i < len(header)} for r in rows[1:]]
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(text)))
    raise ValueError("Upload a .csv or .xlsx file")
